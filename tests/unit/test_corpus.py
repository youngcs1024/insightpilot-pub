"""Offline corpus failures and metric-body drift without model or database dependencies."""

# ruff: noqa: PLR2004 -- small synthetic counts and fixed catalog acceptance values.

import subprocess
import sys
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest
from pydantic import ValidationError
from pypdf import PdfWriter

from app.core.errors import CorpusValidationError
from app.schemas.corpus import (
    CorpusEntry,
    CorpusFormat,
    CorpusManifest,
    CorpusMetadata,
    CorpusStatistics,
)
from app.schemas.metrics import MetricDefinition
from data.corpus_io import (
    corpus_statistics,
    load_corpus,
    markdown_parts,
    parse_yaml,
    read_document,
    validate_predecessors,
)
from data.corpus_metrics import normative_body, validate_metric_memos
from data.seed.metrics_loader import load_catalog
from scripts.corpus_stats import main
from tests.corpus_support import entry, markdown_source, memo, metadata_source, write_inventory

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def one_source(tmp_path: Path) -> Path:
    (tmp_path / "rule.md").write_text(markdown_source(), encoding="utf-8")
    write_inventory(tmp_path, [entry()])
    return tmp_path


@pytest.fixture
def definitions() -> list[MetricDefinition]:
    return load_catalog(ROOT / "data/seed/metrics.yaml").definitions


@pytest.mark.parametrize(
    "name", ["../rule.md", "/rule.md", "a/../rule.md", "a//rule.md", "a\\rule.md"]
)
def test_noncanonical_paths_are_rejected(name: str) -> None:
    with pytest.raises(ValidationError):
        CorpusEntry(path=name, format=CorpusFormat.MARKDOWN, metadata_path=None)


def test_format_and_sidecar_must_match() -> None:
    with pytest.raises(ValidationError):
        CorpusEntry(path="rule.pdf", format=CorpusFormat.MARKDOWN, metadata_path=None)
    with pytest.raises(ValidationError):
        CorpusEntry(path="rule.xlsx", format=CorpusFormat.EXCEL, metadata_path="rule.meta.yaml")


def test_duplicate_manifest_entries_fail() -> None:
    with pytest.raises(ValidationError):
        CorpusManifest(documents=[entry(), entry()])


@pytest.mark.parametrize(
    "source",
    [
        "",
        "title: [",
        "title: one\ntitle: two",
        "value: &recursive [*recursive]",
        metadata_source() + "unknown: true\n",
        metadata_source().replace("policy", "unknown"),
        metadata_source().replace("effective_to: null", "effective_to: 2025-12-31"),
        metadata_source().replace("effective_to: null", "effective_to: 2026-01-01"),
        metadata_source().replace("supersedes: null", "supersedes: ../private.md"),
        metadata_source() + "metric_key: gmv\n",
    ],
)
def test_invalid_yaml_metadata_has_typed_failure(source: str) -> None:
    with pytest.raises(CorpusValidationError):
        parse_yaml(source, CorpusMetadata)


@pytest.mark.parametrize(
    "source", ["# no metadata", "---\ntitle: open", "---\n" + metadata_source() + "---\n"]
)
def test_invalid_front_matter_fails(source: str) -> None:
    with pytest.raises(CorpusValidationError):
        markdown_parts(source)


def test_crlf_normalization_preserves_body() -> None:
    metadata, body = markdown_parts(markdown_source().replace("\n", "\r\n"))
    assert metadata.effective_from == date(2026, 1, 1)
    assert body == "# 测试规则\n\n退款申请需核验商品行。"


def test_missing_file_fails_without_deleting_fixture(one_source: Path) -> None:
    write_inventory(one_source, [entry(), entry("missing.md")])
    with pytest.raises(CorpusValidationError):
        load_corpus(one_source)


def test_extra_file_fails(one_source: Path) -> None:
    (one_source / "unlisted.md").write_text(markdown_source(), encoding="utf-8")
    with pytest.raises(CorpusValidationError):
        load_corpus(one_source)


def test_redteam_corpus_is_excluded_from_normal_inventory() -> None:
    ordinary = load_corpus()
    adversarial = load_corpus(ROOT / "data/corpus/adversarial")
    assert ordinary
    assert len(adversarial) == 3
    assert all(not document.entry.path.startswith("adversarial/") for document in ordinary)


def test_invalid_utf8_fails(one_source: Path) -> None:
    (one_source / "rule.md").write_bytes(b"\xff\xfe\x00")
    with pytest.raises(CorpusValidationError):
        load_corpus(one_source)


def test_symlink_source_fails(tmp_path: Path) -> None:
    (tmp_path / "target.md").write_text(markdown_source(), encoding="utf-8")
    (tmp_path / "rule.md").symlink_to(tmp_path / "target.md")
    write_inventory(tmp_path, [entry(), entry("target.md")])
    with pytest.raises(CorpusValidationError):
        load_corpus(tmp_path)


@pytest.mark.parametrize("name", ["rule.xlsx", "rule.pdf"])
def test_damaged_binary_source_fails(tmp_path: Path, name: str) -> None:
    (tmp_path / name).write_bytes(b"invalid binary document")
    (tmp_path / (name + ".meta.yaml")).write_text(metadata_source(), encoding="utf-8")
    write_inventory(tmp_path, [entry(name)])
    with pytest.raises(CorpusValidationError):
        load_corpus(tmp_path)


def test_zip_without_workbook_fails(tmp_path: Path) -> None:
    with ZipFile(tmp_path / "rule.xlsx", "w") as archive:
        archive.writestr("unrelated.xml", "<unrelated/>")
    (tmp_path / "rule.xlsx.meta.yaml").write_text(metadata_source(), encoding="utf-8")
    with pytest.raises(CorpusValidationError):
        read_document(tmp_path, entry("rule.xlsx"))


def test_missing_binary_sidecar_fails(tmp_path: Path) -> None:
    (tmp_path / "rule.xlsx").write_bytes(b"content")
    write_inventory(tmp_path, [entry("rule.xlsx")])
    with pytest.raises(CorpusValidationError):
        load_corpus(tmp_path)


@pytest.mark.parametrize("encrypted", [False, True])
def test_pdf_without_text_or_with_encryption_fails(tmp_path: Path, encrypted: bool) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    if encrypted:
        writer.encrypt("synthetic-corpus-fixture")
    writer.write(tmp_path / "rule.pdf")
    writer.close()
    (tmp_path / "rule.pdf.meta.yaml").write_text(metadata_source(), encoding="utf-8")
    with pytest.raises(CorpusValidationError):
        read_document(tmp_path, entry("rule.pdf"))


def test_excel_headers_sheets_and_literal_cells(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "费用"
    sheet.append(["项目", "规则"])
    sheet.append(["商品|运费", "分别\n核验"])
    other = book.create_sheet("代码")
    other.append(["SKU", "说明"])
    other.append(["SKU-A1023", "培训样例"])
    book.save(tmp_path / "rule.xlsx")
    book.close()
    (tmp_path / "rule.xlsx.meta.yaml").write_text(metadata_source(), encoding="utf-8")
    loaded = read_document(tmp_path, entry("rule.xlsx"))
    assert loaded.text == (
        "## 费用\n\n| 项目 | 规则 |\n| --- | --- |\n| 商品\\|运费 | 分别<br>核验 |"
        "\n\n## 代码\n\n| SKU | 说明 |\n| --- | --- |\n| SKU-A1023 | 培训样例 |"
    )


def test_empty_workbook_fails(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.save(tmp_path / "rule.xlsx")
    book.close()
    (tmp_path / "rule.xlsx.meta.yaml").write_text(metadata_source(), encoding="utf-8")
    with pytest.raises(CorpusValidationError):
        read_document(tmp_path, entry("rule.xlsx"))


@pytest.mark.parametrize("failure", ["missing", "gap", "cycle", "fork"])
def test_invalid_predecessor_chain_fails(one_source: Path, failure: str) -> None:
    old = load_corpus(one_source)[0]
    old.metadata.effective_to = date(2026, 2, 1)
    new = old.model_copy(deep=True)
    new.entry = entry("new.md")
    new.metadata.effective_from = date(2026, 2, 1)
    new.metadata.effective_to = None
    new.metadata.supersedes = old.entry.path
    documents = [old, new]
    if failure == "missing":
        new.metadata.supersedes = "unknown.md"
    elif failure == "gap":
        old.metadata.effective_to = date(2026, 1, 31)
    elif failure == "cycle":
        old.metadata.supersedes = new.entry.path
    else:
        duplicate = new.model_copy(deep=True)
        duplicate.entry = entry("fork.md")
        documents.append(duplicate)
    with pytest.raises(CorpusValidationError):
        validate_predecessors(documents)


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("默认日期字段\uff1ar.requested_at", "默认日期字段\uff1ao.paid_at"),
        ("r.status <> 'rejected'", "r.status = 'completed'"),
        ("零分母返回 NULL", "零分母返回 0"),
        ("支持粒度\uff1atotal、day、week、month、region", "支持粒度\uff1atotal、category"),
    ],
)
def test_memo_body_drift_fails(
    definitions: list[MetricDefinition], before: str, after: str
) -> None:
    documents = [memo(definition) for definition in definitions]
    document = next(item for item in documents if item.metadata.metric_key == "refund_rate")
    assert before in document.text
    document.text = document.text.replace(before, after)
    with pytest.raises(CorpusValidationError):
        validate_metric_memos(documents, definitions)


def test_missing_or_duplicate_metric_memos_fail(definitions: list[MetricDefinition]) -> None:
    documents = [memo(definition) for definition in definitions]
    validate_metric_memos(documents, definitions)
    with pytest.raises(CorpusValidationError):
        validate_metric_memos(documents[:-1], definitions)
    with pytest.raises(CorpusValidationError):
        validate_metric_memos([*documents, documents[0]], definitions)


@pytest.mark.parametrize(
    "body",
    [
        "只有指标名称 refund_rate",
        "## 使用说明\n\n解释\n\n## 规范定义\n\n定义",
        "## 规范定义\n\n定义\n\n## 规范定义\n\n另一项\n\n## 使用说明\n\n解释",
    ],
)
def test_unbounded_normative_section_fails(body: str) -> None:
    with pytest.raises(CorpusValidationError):
        normative_body(body)


def test_statistics_count_body_not_metadata(one_source: Path) -> None:
    documents = load_corpus(one_source)
    result = corpus_statistics(documents)
    assert result.total == 1
    assert result.min_chars == result.max_chars == result.mean_chars == len(documents[0].text)
    assert result.documents[0].char_count < len(markdown_source())
    with pytest.raises(CorpusValidationError):
        corpus_statistics([])


def test_cli_emits_real_inventory_statistics() -> None:
    result = subprocess.run(
        [sys.executable, "-m", "scripts.corpus_stats"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    summary = CorpusStatistics.model_validate_json(result.stdout)
    assert summary.total == len(summary.documents) == 60
    assert summary.by_format == {
        CorpusFormat.MARKDOWN: 54,
        CorpusFormat.EXCEL: 4,
        CorpusFormat.PDF: 2,
    }
    assert "refund_policy_v3.md" in {item.path for item in summary.documents}


def test_cli_failure_is_nonzero_and_safe(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["--root", str(tmp_path)]) == 1
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == "CORPUS_VALIDATION_ERROR: Knowledge corpus validation failed.\n"
