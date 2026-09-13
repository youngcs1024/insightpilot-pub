"""Shared synchronous corpus parsing; callers choose their execution boundary."""

import unicodedata
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl  # type: ignore[import-untyped]
import yaml  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]
from pydantic import ValidationError
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from app.core.errors import CorpusValidationError, SchemaMetadataError
from app.schemas.corpus import (
    CorpusDocument,
    CorpusEntry,
    CorpusFormat,
    CorpusMetadata,
)
from app.schemas.mcp import Contract
from data.seed.schema_metadata_loader import check_keys


def parse_yaml[Model: Contract](source: str, model: type[Model]) -> Model:
    """Validate safe YAML without losing duplicate keys or accepting recursive aliases."""
    try:
        node = yaml.compose(source)
        if node is None:
            raise CorpusValidationError("Empty YAML metadata.")
        check_keys(node)
        return model.model_validate(yaml.safe_load(source))
    except (yaml.YAMLError, ValidationError, SchemaMetadataError) as exc:
        raise CorpusValidationError("Invalid corpus metadata.") from exc


def source_path(root: Path, relative: str) -> Path:
    """Reject symlinks, including parent directories, rather than following external files."""
    root = root.resolve()
    path = root / relative
    if not path.resolve().is_relative_to(root.resolve()):
        raise CorpusValidationError("Source escapes corpus root.", path=relative)
    current = path
    while current != root:
        if current.is_symlink():
            raise CorpusValidationError("Symlink in corpus inventory.", path=relative)
        current = current.parent
    return path


def read_utf8(path: Path) -> str:
    """Turn resource decoding failures into the corpus error contract."""
    try:
        return path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        raise CorpusValidationError("Cannot read corpus text.", path=path.name) from exc


def markdown_parts(source: str) -> tuple[CorpusMetadata, str]:
    """Require opening and closing front-matter delimiters on their own lines."""
    normalized = source.replace("\r\n", "\n").replace("\r", "\n")
    if not normalized.startswith("---\n"):
        raise CorpusValidationError("Missing front matter.")
    metadata, delimiter, body = normalized[4:].partition("\n---\n")
    if not delimiter or not body.strip():
        raise CorpusValidationError("Missing front matter boundary or body.")
    return parse_yaml(metadata, CorpusMetadata), body.strip()


def excel_text(path: Path) -> str:
    """Expose each sheet as a header-bearing table without implementing table splitting."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows = [
                [str(cell) if cell is not None else "" for cell in row]
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None for cell in row)
            ]
            if rows:
                header, *body = rows
                table = [header, ["---"] * len(header), *body]
                sections.append(
                    "## " + sheet.title + "\n\n" + "\n".join(table_row(row) for row in table)
                )
        return "\n\n".join(sections)
    finally:
        workbook.close()


def table_row(cells: list[str]) -> str:
    """Keep literal separators and embedded newlines inside their original table cell."""
    escaped = [cell.replace("|", "\\|").replace("\n", "<br>") for cell in cells]
    return "| " + " | ".join(escaped) + " |"


def pdf_text(path: Path) -> str:
    """Read real text pages; scanned or encrypted resources are not valid corpus inputs."""
    with path.open("rb") as stream:
        reader = PdfReader(stream, strict=True)
        if reader.is_encrypted:
            raise CorpusValidationError("Encrypted PDFs are unsupported.", path=path.name)
        pages = [page.extract_text() or "" for page in reader.pages]
        if not pages or any(not page.strip() for page in pages):
            raise CorpusValidationError("PDF has a page without extractable text.", path=path.name)
        return "\n\n".join(pages)


def read_document(root: Path, entry: CorpusEntry) -> CorpusDocument:
    """Load one body and its sole business metadata source."""
    path = source_path(root, entry.path)
    try:
        if entry.format is CorpusFormat.MARKDOWN:
            metadata, text = markdown_parts(read_utf8(path))
        else:
            if entry.metadata_path is None:
                raise CorpusValidationError("Missing binary metadata source.")
            metadata = parse_yaml(read_utf8(source_path(root, entry.metadata_path)), CorpusMetadata)
            text = excel_text(path) if entry.format is CorpusFormat.EXCEL else pdf_text(path)
        return CorpusDocument(
            entry=entry, metadata=metadata, text=unicodedata.normalize("NFC", text.strip())
        )
    except (
        OSError,
        BadZipFile,
        PyPdfError,
        InvalidFileException,
        ParseError,
        ValueError,
        KeyError,
        TypeError,
    ) as exc:
        raise CorpusValidationError("Cannot extract corpus document.", path=entry.path) from exc


